# -*- coding: utf-8 -*-
"""
Export-Funktionen für CSV und Excel.
"""

import csv
from .formatting import DeutscheFormatierung


class ExportManager:
    """Exportiert Daten in verschiedene Formate."""

    def __init__(self, daten, analyse):
        """
        Initialisiert den Export-Manager.

        Args:
            daten: Liste von Lieferschein-Dictionaries
            analyse: VerkaufsAnalyse-Instanz
        """
        self.daten = daten
        self.analyse = analyse
        self.fmt = DeutscheFormatierung()

    def export_csv(self, pfad):
        """
        Exportiert Rohdaten als CSV mit deutscher Formatierung.

        Args:
            pfad: Ausgabepfad für die CSV-Datei
        """
        fieldnames = [
            'Lieferschein_Nr', 'Datum', 'Kunde_Nr', 'Amazon_Bestellung',
            'Bestelldatum', 'Betrag_EUR', 'Artikel_Nr', 'Produkt_Typ',
            'Farbe', 'Menge', 'Land_Code', 'Land'
        ]

        with open(pfad, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=';')
            writer.writeheader()

            for d in self.daten:
                row = {k: d.get(k, '') for k in fieldnames}
                # Betrag in deutschem Format
                row['Betrag_EUR'] = self.fmt.zahl(d.get('Betrag', 0), 2)
                writer.writerow(row)

        print(f"CSV exportiert: {pfad}")

    def export_excel(self, pfad):
        """
        Exportiert alle Analysen als Excel-Datei.

        Args:
            pfad: Ausgabepfad für die Excel-Datei
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            print(
                "FEHLER: openpyxl nicht installiert. "
                "Installieren mit: pip install openpyxl"
            )
            return

        wb = Workbook()

        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        title_font = Font(bold=True, size=14)
        subtitle_font = Font(bold=True, size=12)

        # Deutsche Formate für Excel
        euro_format = '#.##0,00 €'
        pct_format = '0,0%'

        def style_header(ws, row=1):
            for cell in ws[row]:
                if cell.value:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center')

        def auto_width(ws):
            for col_idx in range(1, ws.max_column + 1):
                max_length = 0
                for row_idx in range(1, min(ws.max_row + 1, 200)):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[get_column_letter(col_idx)].width = \
                    min(max_length + 3, 45)

        # Hilfsvariablen
        total_count = len(self.daten)
        betraege = [d['Betrag'] for d in self.daten if d.get('Betrag')]
        total_umsatz = sum(betraege)

        # ===== SHEET 1: Rohdaten =====
        ws1 = wb.active
        ws1.title = "Rohdaten"

        headers = [
            'Lieferschein_Nr', 'Datum', 'Kunde_Nr', 'Amazon_Bestellung',
            'Bestelldatum', 'Betrag', 'Artikel_Nr', 'Produkt_Typ',
            'Farbe', 'Menge', 'Land_Code', 'Land'
        ]
        ws1.append(headers)
        style_header(ws1)

        for d in self.daten:
            row = [d.get(h, '') for h in headers]
            ws1.append(row)

        for row in range(2, len(self.daten) + 2):
            ws1.cell(row=row, column=6).number_format = euro_format

        auto_width(ws1)

        # ===== SHEET 2: Zusammenfassung =====
        ws2 = wb.create_sheet("Zusammenfassung")
        zf = self.analyse.zusammenfassung()

        ws2['A1'] = "VERKAUFSANALYSE - Roboterwerk GmbH"
        ws2['A1'].font = Font(bold=True, size=16)

        daten_liste = [d['Datum'] for d in self.daten if d.get('Datum')]
        if daten_liste:
            ws2['A3'] = f"Zeitraum: {min(daten_liste)} - {max(daten_liste)}"

        ws2['A5'] = "KENNZAHLEN"
        ws2['A5'].font = subtitle_font

        kennzahlen = [
            ['Metrik', 'Wert'],
            ['Anzahl Bestellungen', zf['Anzahl_Bestellungen']],
            ['Gesamtumsatz', zf['Gesamtumsatz']],
            ['Durchschnittl. Bestellung', zf['Durchschnitt']],
            ['Minimum', zf['Minimum']],
            ['Maximum', zf['Maximum']],
            ['Median', zf['Median']],
            ['Gesamt Stück', zf['Gesamt_Stueck']],
        ]

        for row in kennzahlen:
            ws2.append(row)

        style_header(ws2, 7)
        for row in range(8, 14):
            ws2.cell(row=row, column=2).number_format = euro_format
        ws2.cell(row=7, column=2).number_format = '0'
        ws2.cell(row=14, column=2).number_format = '0'

        auto_width(ws2)

        # ===== SHEET 3: Produkt-Analyse =====
        ws3 = wb.create_sheet("Produkt-Analyse")
        ws3['A1'] = "ANALYSE NACH PRODUKT-TYP"
        ws3['A1'].font = title_font

        typ_stats = self.analyse.nach_produkt_typ()

        ws3.append([])
        ws3.append([
            'Produkt-Typ', 'Bestellungen', 'Anteil', 'Stück',
            'Umsatz', 'Umsatz-Anteil', 'Ø Bestellung'
        ])
        style_header(ws3, 3)

        row_num = 4
        for typ, stats in sorted(typ_stats.items(), key=lambda x: -x[1]['umsatz']):
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            ws3.append([
                typ, stats['count'], stats['count']/total_count, stats['stueck'],
                stats['umsatz'],
                stats['umsatz']/total_umsatz if total_umsatz else 0, avg
            ])
            ws3.cell(row=row_num, column=3).number_format = pct_format
            ws3.cell(row=row_num, column=5).number_format = euro_format
            ws3.cell(row=row_num, column=6).number_format = pct_format
            ws3.cell(row=row_num, column=7).number_format = euro_format
            row_num += 1

        # Summenzeile
        total_stueck = sum(s['stueck'] for s in typ_stats.values())
        ws3.append([
            'GESAMT', total_count, 1, total_stueck,
            total_umsatz, 1, total_umsatz/total_count
        ])
        for col in range(1, 8):
            ws3.cell(row=row_num, column=col).font = Font(bold=True)
        ws3.cell(row=row_num, column=3).number_format = pct_format
        ws3.cell(row=row_num, column=5).number_format = euro_format
        ws3.cell(row=row_num, column=6).number_format = pct_format
        ws3.cell(row=row_num, column=7).number_format = euro_format

        auto_width(ws3)

        # ===== SHEET 4: Artikel-Details =====
        ws4 = wb.create_sheet("Artikel-Details")
        ws4['A1'] = "DETAILLIERTE ARTIKEL-ANALYSE"
        ws4['A1'].font = title_font

        art_stats = self.analyse.nach_artikel()

        ws4.append([])
        ws4.append([
            'Artikel-Nr', 'Bestellungen', 'Anteil', 'Stück',
            'Umsatz', 'Ø Bestellung'
        ])
        style_header(ws4, 3)

        row_num = 4
        for art, stats in sorted(art_stats.items(), key=lambda x: -x[1]['count']):
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            ws4.append([
                art, stats['count'], stats['count']/total_count,
                stats['stueck'], stats['umsatz'], avg
            ])
            ws4.cell(row=row_num, column=3).number_format = pct_format
            ws4.cell(row=row_num, column=5).number_format = euro_format
            ws4.cell(row=row_num, column=6).number_format = euro_format
            row_num += 1

        auto_width(ws4)

        # ===== SHEET 5: Tages-Analyse =====
        ws5 = wb.create_sheet("Tages-Analyse")
        ws5['A1'] = "UMSATZ PRO TAG"
        ws5['A1'].font = title_font

        tag_stats = self.analyse.nach_tag()

        ws5.append([])
        ws5.append(['Datum', 'Bestellungen', 'Umsatz', 'Ø Bestellung', 'Anteil'])
        style_header(ws5, 3)

        row_num = 4
        for tag in sorted(tag_stats.keys(), reverse=True):
            stats = tag_stats[tag]
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            ws5.append([
                tag, stats['count'], stats['umsatz'], avg,
                stats['umsatz']/total_umsatz if total_umsatz else 0
            ])
            ws5.cell(row=row_num, column=3).number_format = euro_format
            ws5.cell(row=row_num, column=4).number_format = euro_format
            ws5.cell(row=row_num, column=5).number_format = pct_format
            row_num += 1

        auto_width(ws5)

        # ===== SHEET 6: Farben =====
        ws6 = wb.create_sheet("Farben-Analyse")
        ws6['A1'] = "FARBEN-VERTEILUNG"
        ws6['A1'].font = title_font

        farb_stats = self.analyse.nach_farbe()

        ws6.append([])
        ws6.append(['Farbe', 'Anzahl', 'Anteil', 'Umsatz', 'Ø Bestellung'])
        style_header(ws6, 3)

        row_num = 4
        for farbe, stats in sorted(farb_stats.items(), key=lambda x: -x[1]['count']):
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            ws6.append([
                farbe, stats['count'], stats['count']/total_count,
                stats['umsatz'], avg
            ])
            ws6.cell(row=row_num, column=3).number_format = pct_format
            ws6.cell(row=row_num, column=4).number_format = euro_format
            ws6.cell(row=row_num, column=5).number_format = euro_format
            row_num += 1

        auto_width(ws6)

        # ===== SHEET 7: Länder =====
        ws7 = wb.create_sheet("Länder-Analyse")
        ws7['A1'] = "GEOGRAFISCHE VERTEILUNG"
        ws7['A1'].font = title_font

        land_stats = self.analyse.nach_land()

        ws7.append([])
        ws7.append(['Land', 'Bestellungen', 'Anteil', 'Umsatz', 'Ø Bestellung'])
        style_header(ws7, 3)

        row_num = 4
        for land, stats in sorted(land_stats.items(), key=lambda x: -x[1]['count']):
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            ws7.append([
                land, stats['count'], stats['count']/total_count,
                stats['umsatz'], avg
            ])
            ws7.cell(row=row_num, column=3).number_format = pct_format
            ws7.cell(row=row_num, column=4).number_format = euro_format
            ws7.cell(row=row_num, column=5).number_format = euro_format
            row_num += 1

        auto_width(ws7)

        # ===== SHEET 8: Kunden =====
        ws8 = wb.create_sheet("Kunden-Analyse")
        ws8['A1'] = "KUNDEN-ANALYSE"
        ws8['A1'].font = title_font

        kunde_stats = self.analyse.nach_kunde()

        einmalig = len([k for k, v in kunde_stats.items() if v['count'] == 1])
        mehrfach = len([k for k, v in kunde_stats.items() if v['count'] > 1])

        ws8.append([])
        ws8.append(['Einmalige Kunden', einmalig])
        ws8.append(['Mehrfach-Käufer', mehrfach])
        ws8.append(['Gesamt Kunden', len(kunde_stats)])
        ws8.append([])

        ws8.append(['TOP 20 KUNDEN NACH UMSATZ'])
        ws8['A7'].font = subtitle_font
        ws8.append(['Kunde_Nr', 'Bestellungen', 'Umsatz', 'Ø Bestellung', 'Artikel'])
        style_header(ws8, 8)

        row_num = 9
        for kunde, stats in sorted(
            kunde_stats.items(), key=lambda x: -x[1]['umsatz']
        )[:20]:
            avg = stats['umsatz'] / stats['count'] if stats['count'] > 0 else 0
            artikel = ', '.join(set(stats['artikel']))[:50]
            ws8.append([kunde, stats['count'], stats['umsatz'], avg, artikel])
            ws8.cell(row=row_num, column=3).number_format = euro_format
            ws8.cell(row=row_num, column=4).number_format = euro_format
            row_num += 1

        auto_width(ws8)

        # Speichern
        wb.save(pfad)
        print(f"Excel exportiert: {pfad}")
